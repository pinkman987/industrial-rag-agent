from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("docs/project_overall_plan.docx")


def set_east_asia_font(run, font_name: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.bold = bold
    run.font.size = Pt(10)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F2F4F7")
        set_cell_text(cell, header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_text(cells[index], value)

    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def build_document() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("工业设备运维知识库 RAG 与故障诊断 Agent 系统\n项目总计划")
    set_east_asia_font(run)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("面向 AI 大模型应用开发岗位的企业级项目路线图")
    set_east_asia_font(sub_run)
    sub_run.italic = True
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(85, 85, 85)

    doc.add_heading("一、项目总体定位", level=1)
    doc.add_paragraph(
        "本项目面向制造业设备运维场景，目标是构建一个能够处理设备说明书、维修手册、SOP 文档、报警记录和维修工单的 AI 应用系统。"
        "系统先通过 RAG 完成知识库问答，再逐步引入 MySQL 业务数据、Function Calling、故障诊断 Agent、报告生成和前端演示能力，"
        "最终形成一个适合写入简历并用于面试讲解的企业级大模型应用项目。"
    )

    add_table(
        doc,
        ["项目维度", "规划内容"],
        [
            ["项目名称", "工业设备运维知识库 RAG 与故障诊断 Agent 系统"],
            ["目标岗位", "AI 大模型应用开发 / RAG 应用开发 / Agent 应用开发"],
            ["核心场景", "现场工程师遇到设备故障时，快速查询知识库和历史业务数据，生成排查建议和诊断报告。"],
            ["总体路线", "先跑通 RAG 最小闭环，再加入业务数据库，随后实现工具调用、Agent 诊断和前端展示。"],
            ["当前状态", "阶段一已验收通过，阶段二第一批数据库基础能力正在推进。"],
        ],
        [Inches(1.6), Inches(4.9)],
    )

    doc.add_heading("二、项目阶段总览", level=1)
    add_table(
        doc,
        ["阶段", "名称", "主要目标", "状态"],
        [
            ["阶段一", "RAG 最小闭环", "完成文档上传、解析、切块、向量化、Chroma 检索、DeepSeek 回答和 sources 引用返回。", "已验收通过"],
            ["阶段二", "MySQL 业务数据", "加入 Docker MySQL、SQLAlchemy、设备表、报警记录表、维修记录表、对话记录表和初始化脚本。", "开发中"],
            ["阶段三", "Function Calling 工具调用", "让大模型能够调用知识库检索、设备查询、报警查询、维修查询等工具，并返回可观测 tool_calls。", "待开始"],
            ["阶段四", "故障诊断 Agent 与报告生成", "实现自动组合知识库、业务数据和分析逻辑的故障诊断 Agent，并生成 Markdown 诊断报告。", "待开始"],
            ["阶段五", "Streamlit 前端与项目文档", "提供简版演示页面，补齐 README、验收文档、简历写法和面试话术。", "待开始"],
        ],
        [Inches(0.8), Inches(1.35), Inches(3.55), Inches(0.82)],
    )

    stages = [
        (
            "阶段一：RAG 最小闭环",
            "先证明核心 AI 应用链路可运行：企业文档可以进入知识库，用户问题可以检索到依据，大模型可以基于依据回答，并且回答必须带引用来源。",
            [
                "支持 txt、pdf、docx 文档解析。",
                "将文档切分为 chunk，并保留 file_name、page、chunk_index 等元数据。",
                "使用智谱 BigModel embedding-3 生成向量。",
                "使用 Chroma 持久化向量库并执行相似度检索。",
                "使用 DeepSeek Chat 模型生成回答。",
                "返回 answer 和 sources；知识库外问题返回 sources=[]，不编造引用。",
            ],
            "GET /health 正常；POST /api/chat/ask 对 SV300、CONV_A1、CNC_X200 问题能返回 answer 和 sources；知识库外问题不编造引用。",
        ),
        (
            "阶段二：MySQL 业务数据",
            "在 RAG 文档知识库之外，加入结构化业务数据，为后续工具调用和 Agent 诊断打基础。",
            [
                "使用 Docker 启动 MySQL，数据库名 industrial_rag。",
                "使用 SQLAlchemy + PyMySQL，避免 Windows 下 mysqlclient 编译问题。",
                "建设 devices、alarm_records、maintenance_records、chat_records 四张表。",
                "提供 scripts/init_db.py 初始化示例设备、报警记录和维修记录，并支持重复执行。",
                "后续批次新增设备查询接口、对话记录保存和历史查询。",
            ],
            "docker compose up -d mysql 可启动数据库；python scripts/init_db.py 可创建表并插入示例数据；重复执行不产生重复脏数据；阶段一 RAG 问答不受影响。",
        ),
        (
            "阶段三：Function Calling 工具调用",
            "让大模型不仅能生成文本，还能根据问题选择并调用系统工具，增强面试可讲解性和工程完整性。",
            [
                "实现 search_knowledge_base 工具。",
                "实现 query_device_info 工具。",
                "实现 query_alarm_records 工具。",
                "实现 query_maintenance_records 工具。",
                "问答接口返回 tool_calls，记录模型调用了哪些工具、输入参数和输出摘要。",
            ],
            "指定 device_code 提问时能看到数据库工具调用；文档类问题能看到知识库检索工具调用；tool_calls 可用于面试讲解模型决策路径。",
        ),
        (
            "阶段四：故障诊断 Agent 与报告生成",
            "在 RAG、数据库和工具调用稳定后，实现面向设备故障场景的诊断 Agent，输出结构化分析结果。",
            [
                "Agent 根据问题自动组合知识库检索、设备信息查询、报警查询和维修查询。",
                "生成故障现象、知识库依据、历史报警依据、历史维修依据、可能原因、排查步骤、建议措施。",
                "新增 POST /api/reports/generate 报告生成接口。",
                "报告使用 Markdown，便于后续前端展示或导出。",
            ],
            "输入设备故障问题后，Agent 能返回诊断建议；报告中明确区分知识库依据和业务数据依据；没有依据时明确说明依据不足。",
        ),
        (
            "阶段五：Streamlit 前端与项目文档",
            "补齐演示层和项目包装，让系统适合 GitHub 展示、简历投递和面试演示。",
            [
                "提供简版 Streamlit 页面。",
                "支持文档上传、问题输入、设备选择、答案展示、引用来源展示、工具调用展示和报告展示。",
                "完善 README、阶段验收文档、项目架构说明、简历写法和面试话术。",
                "保持前端轻量，不做复杂状态管理和权限系统。",
            ],
            "前端可以完成核心演示流程；README 按步骤可以跑通项目；文档能支撑简历和面试讲解。",
        ),
    ]

    for heading, goal, items, acceptance in stages:
        doc.add_heading(heading, level=1)
        doc.add_heading("阶段目标", level=2)
        doc.add_paragraph(goal)
        doc.add_heading("主要建设内容", level=2)
        for item in items:
            add_bullet(doc, item)
        doc.add_heading("验收重点", level=2)
        doc.add_paragraph(acceptance)

    doc.add_heading("三、总体技术路线", level=1)
    add_table(
        doc,
        ["模块", "技术选择", "作用"],
        [
            ["后端接口", "FastAPI", "提供文档上传、问答、设备查询、历史记录和报告生成等 API。"],
            ["RAG 编排", "LangChain", "连接文档解析、Embedding、向量检索和大模型回答。"],
            ["向量数据库", "Chroma", "持久化文档 chunk 向量并执行相似度检索。"],
            ["Chat 模型", "DeepSeek", "根据检索内容生成设备运维问答和诊断建议。"],
            ["Embedding 模型", "智谱 BigModel embedding-3", "将文档片段和用户问题转换成向量。"],
            ["业务数据库", "MySQL + SQLAlchemy + PyMySQL", "保存设备、报警、维修和对话记录。"],
            ["前端展示", "Streamlit", "用于后期快速搭建项目演示页面。"],
        ],
        [Inches(1.35), Inches(1.8), Inches(3.35)],
    )

    doc.add_heading("四、最终项目能力画像", level=1)
    for item in [
        "能处理企业非结构化文档，构建可追溯的设备运维知识库。",
        "能结合结构化业务数据，查询设备、报警和维修历史。",
        "能通过 Function Calling 展示大模型调用外部工具的过程。",
        "能通过 Agent 组合多种工具完成故障分析和报告生成。",
        "能通过前端页面完成面试演示，体现完整的大模型应用开发能力。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("五、简历表达建议", level=1)
    doc.add_paragraph(
        "可在简历中描述为：设计并实现工业设备运维知识库 RAG 与故障诊断 Agent 系统，基于 FastAPI、LangChain、Chroma、MySQL 和大模型 API，"
        "完成设备文档向量化检索、带引用问答、业务数据建模、工具调用和故障诊断报告生成。项目突出 RAG 工程化、结构化业务数据接入、"
        "Function Calling、Agent 工作流和可演示产品化能力。"
    )

    doc.add_heading("六、面试讲解主线", level=1)
    for text in [
        "先说明业务背景：制造业现场工程师需要快速查询设备文档和历史维修经验。",
        "再说明架构拆分：非结构化文档进入 Chroma，结构化业务数据进入 MySQL。",
        "重点讲阶段一：如何保证回答有引用、无依据时不编造。",
        "继续讲阶段二：为什么引入设备表、报警表、维修表和对话记录表。",
        "最后讲演进方向：Function Calling 让模型调用工具，Agent 负责自动诊断和报告生成。",
    ]:
        add_number(doc, text)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("工业设备运维知识库 RAG 与故障诊断 Agent 系统 - 项目总计划")
    set_east_asia_font(footer_run)
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(85, 85, 85)

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    build_document()

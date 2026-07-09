from pathlib import Path
from typing import Optional


class UnsupportedFileTypeError(ValueError):
    pass


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, "无法用 utf-8/gb18030 解码文本文件")


def load_document(path: Path, file_name: Optional[str] = None):
    """Parse pdf/txt/docx into LangChain Documents with stable metadata."""
    from langchain_core.documents import Document

    path = Path(path)
    display_name = file_name or path.name
    suffix = path.suffix.lower()

    if suffix == ".txt":
        text = _read_text_file(path)
        return [Document(page_content=text, metadata={"file_name": display_name, "page": None})]

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        documents = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(page_content=text, metadata={"file_name": display_name, "page": index})
                )
        return documents

    if suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs)
        return [Document(page_content=text, metadata={"file_name": display_name, "page": None})]

    raise UnsupportedFileTypeError("仅支持 .pdf、.txt、.docx 文档。")
